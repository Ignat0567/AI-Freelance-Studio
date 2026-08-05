from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from design_system import ColorPalette, FontPairing

from .library import get_slide
from .models import SlideSpec

_BLANK_LAYOUT_INDEX = 6  # python-pptx's built-in "Blank" layout
_SLIDE_WIDTH_INCHES = 13.333
_SLIDE_HEIGHT_INCHES = 7.5


def build_presentation(
    selected: Sequence[tuple[str, dict[str, str]]],
    palette: ColorPalette,
    font_pairing: FontPairing,
    destination: Path,
) -> None:
    """Build a real .pptx from a sequence of (slug, content) pairs and save it.

    Every slug must reference a real `SLIDE_LIBRARY` entry; unknown slugs fail
    closed. Content dict keys not declared by that slide's `content_schema` are
    silently dropped; any declared field left unfilled raises, since a blank
    heading/bullet would ship as a visibly broken slide.
    """
    slides_to_build: list[tuple[SlideSpec, dict[str, str]]] = []
    for slug, content in selected:
        spec = get_slide(slug)
        if spec is None:
            raise ValueError(f"unknown_slide_slug:{slug}")
        declared_fields = set(spec.content_fields())
        filtered_content = {key: value for key, value in content.items() if key in declared_fields}
        missing = declared_fields - filtered_content.keys()
        if missing:
            raise ValueError(f"missing_content_fields:{slug}:{','.join(sorted(missing))}")
        slides_to_build.append((spec, filtered_content))

    presentation = Presentation()
    presentation.slide_width = Inches(_SLIDE_WIDTH_INCHES)
    presentation.slide_height = Inches(_SLIDE_HEIGHT_INCHES)
    blank_layout = presentation.slide_layouts[_BLANK_LAYOUT_INDEX]

    for spec, content in slides_to_build:
        slide = presentation.slides.add_slide(blank_layout)
        spec.builder(slide, content, palette, font_pairing)

    destination.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(destination))
