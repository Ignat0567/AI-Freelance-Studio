from __future__ import annotations

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from design_system import ColorPalette, FontPairing

from .._pptx_helpers import hex_to_rgbcolor, set_solid_background
from ..models import SlideSpec


def build(slide, content: dict[str, str], palette: ColorPalette, font_pairing: FontPairing) -> None:
    set_solid_background(slide, palette.background)

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.6))
    title_paragraph = title_box.text_frame.paragraphs[0]
    title_paragraph.text = content.get("headline", "")
    title_paragraph.font.size = Pt(44)
    title_paragraph.font.bold = True
    title_paragraph.font.name = font_pairing.heading_family
    title_paragraph.font.color.rgb = hex_to_rgbcolor(palette.text)
    title_paragraph.alignment = PP_ALIGN.CENTER

    subhead_box = slide.shapes.add_textbox(Inches(1.5), Inches(4.1), Inches(10.3), Inches(1.0))
    subhead_paragraph = subhead_box.text_frame.paragraphs[0]
    subhead_paragraph.text = content.get("subhead", "")
    subhead_paragraph.font.size = Pt(20)
    subhead_paragraph.font.name = font_pairing.body_family
    subhead_paragraph.font.color.rgb = hex_to_rgbcolor(palette.muted)
    subhead_paragraph.alignment = PP_ALIGN.CENTER


TITLE_SLIDE = SlideSpec(
    slug="title",
    display_name="Title Slide",
    description="Full-bleed opening slide with a large centered headline and a supporting subhead.",
    when_to_use="Use as the first slide of any deck, to introduce the company/project/topic.",
    content_schema=(
        ("headline", "The main deck title, short and punchy"),
        ("subhead", "One supporting sentence under the title"),
    ),
    builder=build,
)
