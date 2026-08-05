from __future__ import annotations

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from design_system import ColorPalette, FontPairing

from .._pptx_helpers import hex_to_rgbcolor, set_solid_background
from ..models import SlideSpec


def build(slide, content: dict[str, str], palette: ColorPalette, font_pairing: FontPairing) -> None:
    set_solid_background(slide, palette.background)

    heading_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(1.2))
    heading_paragraph = heading_box.text_frame.paragraphs[0]
    heading_paragraph.text = content.get("heading", "")
    heading_paragraph.font.size = Pt(38)
    heading_paragraph.font.bold = True
    heading_paragraph.font.name = font_pairing.heading_family
    heading_paragraph.font.color.rgb = hex_to_rgbcolor(palette.text)
    heading_paragraph.alignment = PP_ALIGN.CENTER

    message_box = slide.shapes.add_textbox(Inches(1.8), Inches(3.3), Inches(9.7), Inches(1.0))
    message_paragraph = message_box.text_frame.paragraphs[0]
    message_paragraph.text = content.get("message", "")
    message_paragraph.font.size = Pt(18)
    message_paragraph.font.name = font_pairing.body_family
    message_paragraph.font.color.rgb = hex_to_rgbcolor(palette.muted)
    message_paragraph.alignment = PP_ALIGN.CENTER

    button = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.9), Inches(4.7), Inches(3.5), Inches(0.9))
    button.fill.solid()
    button.fill.fore_color.rgb = hex_to_rgbcolor(palette.accent)
    button.line.fill.background()
    button_paragraph = button.text_frame.paragraphs[0]
    button_paragraph.text = content.get("cta_label", "")
    button_paragraph.font.size = Pt(20)
    button_paragraph.font.bold = True
    button_paragraph.font.name = font_pairing.heading_family
    button_paragraph.font.color.rgb = hex_to_rgbcolor(palette.surface)
    button_paragraph.alignment = PP_ALIGN.CENTER


CLOSING_CTA_SLIDE = SlideSpec(
    slug="closing_cta",
    display_name="Closing Call-to-Action Slide",
    description="A closing slide with a heading, a short message, and a filled call-to-action button.",
    when_to_use="Use as the last slide of any deck, to close with a clear next action.",
    content_schema=(
        ("heading", "The closing headline"),
        ("message", "One supporting sentence"),
        ("cta_label", "Short call-to-action button label, e.g. 'Get in touch'"),
    ),
    builder=build,
)
