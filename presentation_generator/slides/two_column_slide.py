from __future__ import annotations

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from design_system import ColorPalette, FontPairing

from .._pptx_helpers import hex_to_rgbcolor, set_solid_background
from ..models import SlideSpec


def _column(slide, left_inches: float, heading_text: str, body_text: str, palette: ColorPalette, font_pairing: FontPairing) -> None:
    heading_box = slide.shapes.add_textbox(Inches(left_inches), Inches(1.8), Inches(5.3), Inches(0.8))
    heading_paragraph = heading_box.text_frame.paragraphs[0]
    heading_paragraph.text = heading_text
    heading_paragraph.font.size = Pt(22)
    heading_paragraph.font.bold = True
    heading_paragraph.font.name = font_pairing.heading_family
    heading_paragraph.font.color.rgb = hex_to_rgbcolor(palette.accent)

    body_box = slide.shapes.add_textbox(Inches(left_inches), Inches(2.6), Inches(5.3), Inches(3.8))
    body_frame = body_box.text_frame
    body_frame.word_wrap = True
    body_paragraph = body_frame.paragraphs[0]
    body_paragraph.text = body_text
    body_paragraph.font.size = Pt(18)
    body_paragraph.font.name = font_pairing.body_family
    body_paragraph.font.color.rgb = hex_to_rgbcolor(palette.text)


def build(slide, content: dict[str, str], palette: ColorPalette, font_pairing: FontPairing) -> None:
    set_solid_background(slide, palette.background)

    heading_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.0))
    heading_paragraph = heading_box.text_frame.paragraphs[0]
    heading_paragraph.text = content.get("heading", "")
    heading_paragraph.font.size = Pt(32)
    heading_paragraph.font.bold = True
    heading_paragraph.font.name = font_pairing.heading_family
    heading_paragraph.font.color.rgb = hex_to_rgbcolor(palette.primary)
    heading_paragraph.alignment = PP_ALIGN.LEFT

    _column(slide, 0.8, content.get("left_heading", ""), content.get("left_body", ""), palette, font_pairing)
    _column(slide, 6.7, content.get("right_heading", ""), content.get("right_body", ""), palette, font_pairing)


TWO_COLUMN_SLIDE = SlideSpec(
    slug="two_column",
    display_name="Two Column Slide",
    description="A heading with two side-by-side columns, each with its own sub-heading and body text.",
    when_to_use="Use for comparisons, before/after, or two related points that deserve equal visual weight.",
    content_schema=(
        ("heading", "The slide heading"),
        ("left_heading", "Left column sub-heading"),
        ("left_body", "Left column body text, 1-2 sentences"),
        ("right_heading", "Right column sub-heading"),
        ("right_body", "Right column body text, 1-2 sentences"),
    ),
    builder=build,
)
