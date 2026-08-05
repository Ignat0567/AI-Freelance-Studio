from __future__ import annotations

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from design_system import ColorPalette, FontPairing

from .._pptx_helpers import hex_to_rgbcolor, set_solid_background
from ..models import SlideSpec

_BULLET_FIELDS = ("bullet_1", "bullet_2", "bullet_3")


def build(slide, content: dict[str, str], palette: ColorPalette, font_pairing: FontPairing) -> None:
    set_solid_background(slide, palette.background)

    heading_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.0))
    heading_paragraph = heading_box.text_frame.paragraphs[0]
    heading_paragraph.text = content.get("heading", "")
    heading_paragraph.font.size = Pt(32)
    heading_paragraph.font.bold = True
    heading_paragraph.font.name = font_pairing.heading_family
    heading_paragraph.font.color.rgb = hex_to_rgbcolor(palette.primary)
    heading_paragraph.alignment = PP_ALIGN.LEFT

    body_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(11.0), Inches(4.5))
    body_frame = body_box.text_frame
    body_frame.word_wrap = True
    for index, field_name in enumerate(_BULLET_FIELDS):
        paragraph = body_frame.paragraphs[0] if index == 0 else body_frame.add_paragraph()
        paragraph.text = f"•  {content.get(field_name, '')}"
        paragraph.font.size = Pt(22)
        paragraph.font.name = font_pairing.body_family
        paragraph.font.color.rgb = hex_to_rgbcolor(palette.text)
        paragraph.space_after = Pt(18)


BULLETS_SLIDE = SlideSpec(
    slug="bullets",
    display_name="Bullet Points Slide",
    description="A heading with exactly three short bullet points below it.",
    when_to_use="Use for an agenda, feature list, or key-points slide.",
    content_schema=(
        ("heading", "The slide heading"),
        ("bullet_1", "First bullet point, one short sentence"),
        ("bullet_2", "Second bullet point, one short sentence"),
        ("bullet_3", "Third bullet point, one short sentence"),
    ),
    builder=build,
)
