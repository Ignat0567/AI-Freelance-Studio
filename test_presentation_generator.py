from __future__ import annotations

import json

import pytest
from pptx import Presentation

from design_system import DEFAULT_TOKENS
from presentation_generator import SLIDE_LIBRARY, build_presentation, select_presentation_slides
from presentation_generator.select_slides import _selection_prompt

pytestmark = pytest.mark.unit


def _fill_all_fields(slide) -> dict[str, str]:
    return {field: f"Sample {field}" for field in slide.content_fields()}


def test_slide_library_entries_are_well_formed():
    slugs = [slide.slug for slide in SLIDE_LIBRARY]
    assert len(slugs) == len(set(slugs))
    for slide in SLIDE_LIBRARY:
        assert slide.content_schema
        field_names = slide.content_fields()
        assert len(field_names) == len(set(field_names))
        assert callable(slide.builder)


def test_build_presentation_writes_a_real_pptx_with_expected_slide_count_and_text(tmp_path):
    selected = [(slide.slug, _fill_all_fields(slide)) for slide in SLIDE_LIBRARY]
    destination = tmp_path / "deck.pptx"

    build_presentation(selected, DEFAULT_TOKENS.light, DEFAULT_TOKENS.font_pairing, destination)

    assert destination.is_file()
    reopened = Presentation(str(destination))
    assert len(reopened.slides._sldIdLst) == len(SLIDE_LIBRARY)
    all_text = "\n".join(shape.text_frame.text for slide in reopened.slides for shape in slide.shapes if shape.has_text_frame)
    assert "Sample headline" in all_text
    assert "Sample cta_label" in all_text


def test_build_presentation_rejects_unknown_slug(tmp_path):
    with pytest.raises(ValueError, match="unknown_slide_slug"):
        build_presentation([("not_a_real_slide", {})], DEFAULT_TOKENS.light, DEFAULT_TOKENS.font_pairing, tmp_path / "deck.pptx")


def test_build_presentation_requires_every_declared_field(tmp_path):
    slide = SLIDE_LIBRARY[0]
    content = _fill_all_fields(slide)
    content.pop(slide.content_fields()[0])

    with pytest.raises(ValueError, match="missing_content_fields"):
        build_presentation([(slide.slug, content)], DEFAULT_TOKENS.light, DEFAULT_TOKENS.font_pairing, tmp_path / "deck.pptx")


def test_build_presentation_drops_undeclared_content_keys(tmp_path):
    slide = SLIDE_LIBRARY[0]
    content = {**_fill_all_fields(slide), "not_a_real_field": "should be dropped"}
    destination = tmp_path / "deck.pptx"

    build_presentation([(slide.slug, content)], DEFAULT_TOKENS.light, DEFAULT_TOKENS.font_pairing, destination)

    reopened = Presentation(str(destination))
    all_text = "\n".join(shape.text_frame.text for s in reopened.slides for shape in s.shapes if shape.has_text_frame)
    assert "should be dropped" not in all_text


def _fake_ai_ask_returning(payload: dict):
    def _ask(_prompt: str) -> str:
        return json.dumps(payload)

    return _ask


def test_select_presentation_slides_returns_validated_selection():
    payload = {
        "slides": [
            {"slug": "title", "content": {"headline": "AI Freelance Studio", "subhead": "Built by its own agents"}},
        ]
    }
    ai_ask = _fake_ai_ask_returning(payload)

    selected = select_presentation_slides("AI Freelance Studio pitch deck", ai_ask)

    assert len(selected) == 1
    assert selected[0].slug == "title"
    assert selected[0].content == {"headline": "AI Freelance Studio", "subhead": "Built by its own agents"}


def test_select_presentation_slides_drops_unknown_slug_and_undeclared_fields():
    payload = {
        "slides": [
            {"slug": "not_a_real_slide", "content": {"headline": "should be dropped entirely"}},
            {"slug": "title", "content": {"headline": "Real Headline", "subhead": "Real Subhead", "made_up_field": "dropped"}},
        ]
    }
    ai_ask = _fake_ai_ask_returning(payload)

    selected = select_presentation_slides("topic", ai_ask)

    assert len(selected) == 1
    assert selected[0].slug == "title"
    assert selected[0].content == {"headline": "Real Headline", "subhead": "Real Subhead"}


def test_select_presentation_slides_raises_when_ai_returns_no_valid_slides():
    ai_ask = _fake_ai_ask_returning({"slides": [{"slug": "not_real", "content": {}}]})

    with pytest.raises(ValueError, match="ai_returned_no_valid_slides"):
        select_presentation_slides("topic", ai_ask)


def test_select_presentation_slides_raises_on_invalid_json():
    with pytest.raises(ValueError, match="ai_response_not_valid_json"):
        select_presentation_slides("topic", lambda _prompt: "not json at all")


def test_selection_prompt_lists_every_library_slug():
    prompt = _selection_prompt("some topic")
    for slide in SLIDE_LIBRARY:
        assert slide.slug in prompt
