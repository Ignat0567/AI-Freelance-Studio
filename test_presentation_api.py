from __future__ import annotations

import json
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation

from api.presentation import install_presentation_api, router
from backend_security import LocalSecurityContext, LocalSecurityMiddleware, set_app_security_context
from presentation_generator import SLIDE_LIBRARY

pytestmark = pytest.mark.unit
TOKEN = "presentation-api-focused-token-32-bytes"
ORIGIN = "http://127.0.0.1:8080"


def _fake_ai_ask_selecting_every_slide(_prompt: str) -> str:
    payload = {
        "slides": [
            {"slug": slide.slug, "content": {field: f"Sample {field}" for field in slide.content_fields()}}
            for slide in SLIDE_LIBRARY
        ]
    }
    return json.dumps(payload)


def _fake_ai_ask_returning_garbage(_prompt: str) -> str:
    return "not json at all"


def _app(*, ai_ask=None):
    app = FastAPI()
    app.add_middleware(LocalSecurityMiddleware)
    set_app_security_context(
        app,
        LocalSecurityContext.create(token=TOKEN, bind_host="127.0.0.1", port=8080, launch_id="presentation-api-test-launch", allow_test_client=True),
    )
    install_presentation_api(app, ai_ask=ai_ask)
    app.include_router(router)
    return app


def _client(**kwargs) -> TestClient:
    headers = {"X-FreelancerStudio-Token": TOKEN, "Origin": ORIGIN}
    return TestClient(_app(**kwargs), base_url=ORIGIN, headers=headers)


def test_generate_returns_a_real_valid_pptx():
    response = _client(ai_ask=_fake_ai_ask_selecting_every_slide).post("/api/presentation/generate", json={"topic": "AI Freelance Studio"})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml.presentation")
    presentation = Presentation(BytesIO(response.content))
    assert len(presentation.slides._sldIdLst) == len(SLIDE_LIBRARY)
    all_text = "\n".join(shape.text_frame.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)
    assert "Sample headline" in all_text


def test_generate_returns_502_when_ai_response_is_unusable():
    response = _client(ai_ask=_fake_ai_ask_returning_garbage).post("/api/presentation/generate", json={"topic": "AI Freelance Studio"})

    assert response.status_code == 502
    assert response.json()["detail"]["code"] == "presentation_generation_failed"


def test_generate_requires_the_security_token():
    unauthenticated = TestClient(_app(ai_ask=_fake_ai_ask_selecting_every_slide), base_url=ORIGIN)

    response = unauthenticated.post("/api/presentation/generate", json={"topic": "AI Freelance Studio"})

    assert response.status_code in (401, 403)
